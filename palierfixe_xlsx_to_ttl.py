#!/usr/bin/env python3
"""One-off: PalierFixeMapping *.xlsx → PalierFixeElement.ttl (PF-* sheets only)."""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl

PFIX_NS = "https://codelist.commonapproach.org/PalierFixeElement#"
PFIX_BASE = "https://codelist.commonapproach.org/PalierFixeElement"
POPSERV_NS = "https://codelist.commonapproach.org/PopulationServed#"
CIDS_NS = "https://ontology.commonapproach.org/cids#"
PALIER_FIXE_ORG = "https://palierfixe.quebec/Organization/PalierFixe"
SDG_BASE = "https://metadata.un.org/sdg/"

SHEETS = [
    ("PF-Theme", "Themes", "Theme", "cids:Theme"),
    ("PF-Outcome", "Outcomes", "Outcome", "cids:Outcome"),
    ("PF-Indicator", "Indicators", "Indicator", "cids:Indicator"),
    ("PF-Characteristic", "Characteristics", "Characteristic", "cids:Characteristic"),
]


def sheet_rows(wb, name: str) -> list[dict]:
    ws = wb[name]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rows.append({h: row[i] for i, h in enumerate(headers) if h})
    return rows


def split_values(value) -> list[str]:
    if value is None or value == "":
        return []
    return [part.strip() for part in str(value).split(",") if part.strip()]


def turtle_string(value: str) -> str:
    return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'


def is_ncname_start_char(ch: str) -> bool:
    return bool(ch) and (ch.isalpha() or ch == "_")


def curie(prefix: str, local: str, full_uri: str) -> str:
    if local and is_ncname_start_char(local[0]):
        return f"{prefix}:{local}"
    return f"<{full_uri}>"


def fragment(uri: str) -> str | None:
    if "#" in uri:
        return uri.rsplit("#", 1)[-1]
    return None


def uri_to_term(uri: str, by_fragment: dict[str, str]) -> str:
    uri = uri.strip()
    if uri == PALIER_FIXE_ORG:
        return f"<{uri}>"

    frag = fragment(uri)
    if uri.startswith(PFIX_NS) or (frag and uri.startswith(PFIX_BASE)):
        local = by_fragment.get(frag, frag) if frag else None
        if local:
            return curie("pfix", local, f"{PFIX_NS}{local}")
        return f"<{uri}>"

    if uri.startswith(POPSERV_NS) or "/PopulationServed#" in uri:
        local = frag or ""
        return curie("popsv", local, uri)

    if uri.startswith(CIDS_NS) or "/cids#" in uri:
        local = frag or ""
        return curie("cids", local, uri)

    if uri.startswith(SDG_BASE):
        return f"<{uri}>"

    return f"<{uri}>"


def build_fragment_index(*row_sets: list[dict]) -> dict[str, str]:
    index: dict[str, str] = {}
    for rows in row_sets:
        for row in rows:
            ident = row.get("identifier")
            uri = row.get("id")
            if ident:
                index[str(ident)] = str(ident)
            if uri and (frag := fragment(str(uri))):
                if ident:
                    index[frag] = str(ident)
                elif frag not in index:
                    index[frag] = frag
    return index


def emit_header() -> list[str]:
    return [
        "@prefix pfix: <https://codelist.commonapproach.org/PalierFixeElement#> .",
        "@prefix popsv: <https://codelist.commonapproach.org/PopulationServed#> .",
        "@prefix cids: <https://ontology.commonapproach.org/cids#> .",
        "@prefix dcterms: <http://purl.org/dc/terms/> .",
        "@prefix i72: <http://ontology.eil.utoronto.ca/ISO21972/iso21972#> .",
        "@prefix org: <http://ontology.eil.utoronto.ca/tove/organization#> .",
        "@prefix rdf: <http://www.w3.org/1999/02/22-rdf-syntax-ns#> .",
        "@prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .",
        "@prefix owl: <http://www.w3.org/2002/07/owl#> .",
        "@prefix skos: <http://www.w3.org/2004/02/skos/core#> .",
        "@prefix voaf: <http://purl.org/vocommons/voaf#> .",
        "@prefix xsd: <http://www.w3.org/2001/XMLSchema#> .",
        "",
        "#################################################################",
        "# Palier Fixe Elements",
        "#",
        "# Generated from PalierFixeMapping Excel (PF-* sheets).",
        "#################################################################",
        "",
        f"<{PFIX_BASE}>",
        "    a voaf:Vocabulary, owl:Ontology, skos:ConceptScheme ;",
        "    owl:imports <https://ontology.commonapproach.org/cids> ;",
        f"    dcterms:creator <{PALIER_FIXE_ORG}> ;",
        '    dcterms:date "2026-05-28"^^xsd:date ;',
        '    owl:versionInfo "1.2" ;',
        '    skos:prefLabel "Palier Fixe Elements"@en ;',
        '    skos:prefLabel "Palier Fixe Elements"@fr ;',
        '    dcterms:title "Palier Fixe Elements"@en ;',
        '    dcterms:title "Palier Fixe Elements"@fr ;',
        '    dcterms:description "A codelist vocabulary of Palier Fixe measurement terms for use in CIDS-aligned software."@en .',
        "",
    ]


def subject_line(local: str) -> str:
    return curie("pfix", local, f"{PFIX_NS}{local}")


def common_properties(local: str, name: str, cids_type: str) -> list[str]:
    ident_lit = turtle_string(local)
    name_lit = turtle_string(name)
    label_lit = turtle_string(name)
    return [
        f"    a {cids_type}, cids:Code, skos:Concept ;",
        f"    skos:inScheme <{PFIX_BASE}> ;",
        f"    rdfs:isDefinedBy <{PFIX_BASE}> ;",
        f"    org:hasIdentifier {ident_lit} ;",
        f"    skos:notation {ident_lit} ;",
        f"    org:hasName {name_lit} ;",
        f"    skos:prefLabel {label_lit}@fr ;",
        f"    cids:definedBy <{PALIER_FIXE_ORG}> ;",
    ]


def emit_theme(row: dict, by_fragment: dict[str, str]) -> list[str]:
    local = str(row["identifier"])
    lines = [subject_line(local), *common_properties(local, str(row["hasName"]), "cids:Theme")]
    parent = row.get("relatesToId") or (
        fragment(str(row["relatesTo"])) if row.get("relatesTo") else None
    )
    if parent:
        parent = str(parent).strip()
        lines.append(f"    cids:relatesTo {subject_line(parent)} ;")
        lines.append(f"    skos:broader {subject_line(parent)} .")
    else:
        lines[-1] = lines[-1].rstrip(" ;") + " ."
    return lines


def emit_outcome(row: dict, by_fragment: dict[str, str]) -> list[str]:
    local = str(row["identifier"])
    lines = [subject_line(local), *common_properties(local, str(row["hasName"]), "cids:Outcome")]
    desc = row.get("hasDescription")
    if desc:
        d = turtle_string(str(desc))
        lines.append(f"    cids:hasDescription {d}@fr ;")
        lines.append(f"    skos:definition {d}@fr ;")

    themes = [uri_to_term(u, by_fragment) for u in split_values(row.get("forTheme"))]
    if themes:
        lines.append(f"    cids:forTheme {', '.join(themes)} ;")

    org = row.get("forOrganization")
    if org:
        lines.append(f"    cids:forOrganization <{org}> ;")

    indicators = [
        uri_to_term(u, by_fragment) for u in split_values(row.get("hasIndicator"))
    ]
    if indicators:
        lines.append(f"    cids:hasIndicator {indicators[0]},")
        lines.extend(f"        {ind}," for ind in indicators[1:-1])
        if len(indicators) > 1:
            lines.append(f"        {indicators[-1]} .")
        else:
            lines[-1] = lines[-1].rstrip(",") + " ."
        return lines

    lines[-1] = lines[-1].rstrip(" ;") + " ."
    return lines


def emit_indicator(row: dict, by_fragment: dict[str, str]) -> list[str]:
    local = str(row["identifier"])
    name = row.get("hasName (Field - FR)") or row.get("hasName") or ""
    lines = [subject_line(local), *common_properties(local, str(name), "cids:Indicator")]

    desc = row.get("hasDescription")
    if desc:
        lines.append(f"    cids:hasDescription {turtle_string(str(desc))}@fr ;")

    unit = row.get("unit_of_measure")
    if unit:
        lines.append(f"    i72:unit_of_measure {uri_to_term(str(unit), by_fragment)} ;")

    unit_desc = row.get("unitDescription")
    if unit_desc:
        lines.append(f"    cids:unitDescription {turtle_string(str(unit_desc))} ;")

    org = row.get("forOrganization")
    if org:
        lines.append(f"    cids:forOrganization <{org}> ;")

    outcome = row.get("forOutcome")
    if outcome:
        lines.append(f"    cids:forOutcome {uri_to_term(str(outcome), by_fragment)} ;")

    theme = row.get("forTheme")
    if theme:
        lines.append(f"    cids:forTheme {uri_to_term(str(theme), by_fragment)} .")
    else:
        lines[-1] = lines[-1].rstrip(" ;") + " ."
    return lines


def emit_characteristic(row: dict, by_fragment: dict[str, str]) -> list[str]:
    local = str(row["identifier"])
    lines = [
        subject_line(local),
        *common_properties(local, str(row["hasName"]), "cids:Characteristic"),
    ]
    code = row.get("hasCode")
    if code:
        lines.append(f"    cids:hasCode {uri_to_term(str(code), by_fragment)} .")
    else:
        lines[-1] = lines[-1].rstrip(" ;") + " ."
    return lines


EMITTERS = {
    "PF-Theme": emit_theme,
    "PF-Outcome": emit_outcome,
    "PF-Indicator": emit_indicator,
    "PF-Characteristic": emit_characteristic,
}


def convert(xlsx_path: Path, out_path: Path) -> None:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    data = {name: sheet_rows(wb, name) for name, _, _, _ in SHEETS}
    by_fragment = build_fragment_index(*data.values())

    out: list[str] = emit_header()
    for sheet_name, section_title, _type_val, _cids_type in SHEETS:
        out.append("#################################################################")
        out.append(f"#    {section_title}")
        out.append("#################################################################")
        out.append("")
        emit = EMITTERS[sheet_name]
        for row in data[sheet_name]:
            block = emit(row, by_fragment)
            out.extend(block)
            out.append("")

    out_path.write_text("\n".join(out).rstrip() + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert PalierFixeMapping xlsx to TTL.")
    parser.add_argument(
        "xlsx",
        nargs="?",
        default="PalierFixeMapping20260528.xlsx",
        help="Input workbook (default: PalierFixeMapping20260528.xlsx)",
    )
    parser.add_argument(
        "-o",
        "--output",
        default="PalierFixeElement.ttl",
        help="Output Turtle file (default: PalierFixeElement.ttl)",
    )
    args = parser.parse_args()
    xlsx = Path(args.xlsx)
    out = Path(args.output)
    convert(xlsx, out)
    print(f"Wrote {out} ({out.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
